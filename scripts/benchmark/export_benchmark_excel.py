from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side,
)
from openpyxl.utils import get_column_letter

DARK_BLUE  = "1F3864"
MID_BLUE   = "2E75B6"
LIGHT_BLUE = "DDEEFF"
LIGHT_GRAY = "F2F2F2"
GREEN_BG   = "E2EFDA"
RED_BG     = "FCE4D6"
YELLOW_BG  = "FFF2CC"
WHITE      = "FFFFFF"

def _hdr(ws, row, col, text, bg=DARK_BLUE, fg=WHITE, bold=True, wrap=False):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(bold=bold, color=fg, size=10)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=wrap)
    _border(cell)
    return cell


def _cell(ws, row, col, value, bg=None, bold=False, fmt=None, align="right"):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, size=10)
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        cell.number_format = fmt
    _border(cell)
    return cell


def _border(cell):
    thin = Side(style="thin", color="CCCCCC")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _autofit(ws, min_w=8, max_w=40):
    for col_cells in ws.columns:
        length = max(
            (len(str(c.value)) if c.value is not None else 0) for c in col_cells
        )
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, min_w), max_w)


def _fmt_s(s: float | None) -> str:
    if s is None:
        return "N/D"
    if s < 0.001:
        return f"{s*1000:.1f}ms"
    if s < 1:
        return f"{s*1000:.0f}ms"
    return f"{s:.2f}s"


LOAD_RUNS = [
    {"run": 1, "motor": "PostgreSQL", "workers": 1, "modo": "serial",
     "tiempo_min": 33.9, "tiempo_s": 2035, "throughput": 778,
     "speedup": "baseline", "filas": 1584534, "skipped": 27},
    {"run": 2, "motor": "PostgreSQL", "workers": 4, "modo": "paralelo",
     "tiempo_min": 25.6, "tiempo_s": 1535, "throughput": 1032,
     "speedup": "1.3x", "filas": 1584534, "skipped": 27},
    {"run": 3, "motor": "PostgreSQL", "workers": 8, "modo": "paralelo",
     "tiempo_min": 10.9, "tiempo_s": 651, "throughput": 2434,
     "speedup": "3.1x", "filas": 1584534, "skipped": 27},
    {"run": 4, "motor": "Iceberg", "workers": 1, "modo": "serial",
     "tiempo_min": 28.7, "tiempo_s": 1722, "throughput": 917,
     "speedup": "baseline", "filas": 1584561, "skipped": 0},
    {"run": 5, "motor": "Iceberg", "workers": 8, "modo": "paralelo",
     "tiempo_min": 7.4, "tiempo_s": 441, "throughput": 3593,
     "speedup": "3.9x", "filas": 1584561, "skipped": 0},
]

QUERY_DESCS = {
    "Q-A1": "Lookup hacienda por codigopersona",
    "Q-A2": "Lookup sfp por documento",
    "Q-A3": "Lookup set por ruc (sin idx)",
    "Q-B1": "Snapshot mensual hacienda (1 mes)",
    "Q-B2": "Snapshot mensual sfp (1 mes)",
    "Q-B3": "Búsqueda nombre parcial (LIKE)",
    "Q-C1": "Masa salarial por entidad (GROUP BY)",
    "Q-C2": "Empleados distintos por mes (COUNT DISTINCT)",
    "Q-C3": "JOIN hacienda × sfp (LIMIT 1000)",
    "Q-D1": "Conteo total por fuente (UNION ALL 5 tablas)",
    "Q-D2": "Perfil cruzado una persona (UNION ALL filtrado)",
    "Q-E1": "Lookup por raw_data_id (UNIQUE idx)",
    "Q-E2": "Lookup por charge exacto (40K filas)",
    "Q-E3": "Rango fecha 1 año, ~205K filas, LIMIT 500",
    "Q-F1": "SUM/AVG activo+neto por año (GROUP BY)",
    "Q-F2": "Top 20 cargos por activo promedio",
    "Q-F3": "COUNT filas activo en [10M, 100M]",
    "Q-F4": "COUNT(*) total",
    "Q-G1": "PK lookup + extracción parsed->>'totalActivo'",
    "Q-G2": "GROUP BY año desde parsed->>'anioDeclaracion'",
    "Q-G3": "Filtro JSONB @> {tipoPresentacion:DIGITAL}",
}

SCENARIOS_ORDER = ["PG-idx", "PG-noidx", "Iceberg"]


def load_csv(path: Path) -> list[dict]:
    rows = []
    with open(path, newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows.append({
                "query": row["query"],
                "group": row["group"],
                "scenario": row["scenario"],
                "min_s":  float(row["min_s"]),
                "p50_s":  float(row["p50_s"]),
                "p95_s":  float(row["p95_s"]),
                "p99_s":  float(row["p99_s"]),
                "max_s":  float(row["max_s"]),
                "runs":   int(row["runs"]),
            })
    return rows


def merge_e3(results_dir: Path) -> list[dict]:
    rows: dict[tuple, dict] = {}
    for fname in ["e3_djbr_pg.csv", "e3_djbr_full.csv"]:
        p = results_dir / fname
        if not p.exists():
            continue
        for r in load_csv(p):
            key = (r["query"], r["scenario"])
            if key not in rows:
                rows[key] = r
    return list(rows.values())


def sheet_carga(wb: Workbook):
    ws = wb.create_sheet("Carga DJBR (E1)")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    title = ws.cell(row=1, column=1, value="Benchmark de Carga — JSON Oficial DJBR  (1,584,561 archivos)")
    title.font = Font(bold=True, size=12, color=DARK_BLUE)
    ws.merge_cells("A1:J1")
    title.alignment = Alignment(horizontal="center")

    headers = ["Run", "Motor", "Workers", "Modo", "Tiempo (min)",
               "Tiempo (s)", "Throughput (rows/s)", "Speedup vs serial", "Filas", "Skipped"]
    for c, h in enumerate(headers, 1):
        _hdr(ws, 2, c, h, bg=MID_BLUE)

    for i, run in enumerate(LOAD_RUNS, 3):
        bg = LIGHT_BLUE if i % 2 == 0 else WHITE
        is_best = run["throughput"] == max(r["throughput"] for r in LOAD_RUNS)
        row_bg = GREEN_BG if is_best else bg
        _cell(ws, i, 1,  run["run"],           bg=row_bg, align="center")
        _cell(ws, i, 2,  run["motor"],          bg=row_bg, bold=run["motor"]=="Iceberg", align="left")
        _cell(ws, i, 3,  run["workers"],        bg=row_bg, align="center")
        _cell(ws, i, 4,  run["modo"],           bg=row_bg, align="center")
        _cell(ws, i, 5,  run["tiempo_min"],     bg=row_bg, fmt="0.0")
        _cell(ws, i, 6,  run["tiempo_s"],       bg=row_bg, fmt="#,##0")
        _cell(ws, i, 7,  run["throughput"],     bg=row_bg, fmt="#,##0", bold=is_best)
        _cell(ws, i, 8,  run["speedup"],        bg=row_bg, align="center", bold=is_best)
        _cell(ws, i, 9,  run["filas"],          bg=row_bg, fmt="#,##0")
        _cell(ws, i, 10, run["skipped"],        bg=row_bg, align="center")

    note_row = len(LOAD_RUNS) + 4
    ws.cell(row=note_row, column=1,
            value="Fase Iceberg serial: Fase 1 JSONL 27.8min + Fase 2 Spark 53.3s "
                  "| Fase Iceberg paralelo: Fase 1 JSONL 6.4min + Fase 2 Spark 58.7s")
    ws.cell(row=note_row, column=1).font = Font(italic=True, size=9, color="666666")
    ws.merge_cells(f"A{note_row}:J{note_row}")

    _autofit(ws)
    ws.row_dimensions[2].height = 28


def sheet_raw(wb: Workbook, title: str, rows: list[dict]):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    hdr_title = ws.cell(row=1, column=1, value=title)
    hdr_title.font = Font(bold=True, size=12, color=DARK_BLUE)
    ws.merge_cells("A1:J1")
    hdr_title.alignment = Alignment(horizontal="center")

    headers = ["Query", "Grupo", "Escenario", "min (s)", "p50 (s)", "p95 (s)",
               "p99 (s)", "max (s)", "Runs", "Descripción"]
    for c, h in enumerate(headers, 1):
        _hdr(ws, 2, c, h, bg=MID_BLUE)

    sorted_rows = sorted(rows, key=lambda r: (r["group"], r["query"],
                          SCENARIOS_ORDER.index(r["scenario"])
                          if r["scenario"] in SCENARIOS_ORDER else 99))

    for i, r in enumerate(sorted_rows, 3):
        bg = LIGHT_BLUE if i % 2 == 0 else WHITE
        _cell(ws, i, 1,  r["query"],    bg=bg, bold=True, align="left")
        _cell(ws, i, 2,  r["group"],    bg=bg, align="center")
        _cell(ws, i, 3,  r["scenario"], bg=bg, align="left")
        _cell(ws, i, 4,  r["min_s"],    bg=bg, fmt="0.0000")
        _cell(ws, i, 5,  r["p50_s"],    bg=bg, fmt="0.0000", bold=True)
        _cell(ws, i, 6,  r["p95_s"],    bg=bg, fmt="0.0000")
        _cell(ws, i, 7,  r["p99_s"],    bg=bg, fmt="0.0000")
        _cell(ws, i, 8,  r["max_s"],    bg=bg, fmt="0.0000")
        _cell(ws, i, 9,  r["runs"],     bg=bg, align="center")
        _cell(ws, i, 10, QUERY_DESCS.get(r["query"], ""), bg=bg, align="left")

    _autofit(ws)
    ws.row_dimensions[2].height = 28


def sheet_pivot(wb: Workbook, title: str, rows: list[dict], scenarios: list[str]):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C3"

    hdr_title = ws.cell(row=1, column=1, value=title)
    hdr_title.font = Font(bold=True, size=12, color=DARK_BLUE)
    ncols = 2 + len(scenarios) * 3 + 1
    ws.merge_cells(f"A1:{get_column_letter(ncols)}1")
    hdr_title.alignment = Alignment(horizontal="center")

    _hdr(ws, 2, 1, "Query",  bg=DARK_BLUE)
    _hdr(ws, 2, 2, "Descripción", bg=DARK_BLUE, wrap=True)
    col = 3
    sc_colors = {"PG-idx": "1F497D", "PG-noidx": "7F7F7F", "Iceberg": "375623"}
    for sc in scenarios:
        bg = sc_colors.get(sc, MID_BLUE)
        _hdr(ws, 2, col,   f"{sc} p50", bg=bg, wrap=True)
        _hdr(ws, 2, col+1, f"{sc} p95", bg=bg, wrap=True)
        _hdr(ws, 2, col+2, f"{sc} p99", bg=bg, wrap=True)
        col += 3
    _hdr(ws, 2, col, "Ganador (p50)", bg=DARK_BLUE)

    pivot: dict[str, dict] = {}
    for r in rows:
        if r["query"] not in pivot:
            pivot[r["query"]] = {}
        pivot[r["query"]][r["scenario"]] = r

    queries_sorted = sorted(pivot.keys(),
                            key=lambda q: (q.split("-")[1][0], q))

    prev_group = None
    data_row = 3
    for qid in queries_sorted:
        grp = qid.split("-")[1][0]
        if grp != prev_group:
            sep = ws.cell(row=data_row, column=1,
                          value=f"── Grupo {grp} ──────────────────────────────────────")
            sep.font = Font(bold=True, color=DARK_BLUE, size=10)
            sep.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            ws.merge_cells(f"A{data_row}:{get_column_letter(ncols)}{data_row}")
            data_row += 1
            prev_group = grp

        bg = LIGHT_GRAY if data_row % 2 == 0 else WHITE
        sc_data = pivot[qid]

        best_sc = None
        best_p50 = None
        for sc in scenarios:
            if sc in sc_data:
                p = sc_data[sc]["p50_s"]
                if best_p50 is None or p < best_p50:
                    best_p50 = p
                    best_sc = sc

        _cell(ws, data_row, 1, qid, bg=bg, bold=True, align="left")
        _cell(ws, data_row, 2, QUERY_DESCS.get(qid, ""), bg=bg, align="left")

        col = 3
        for sc in scenarios:
            if sc in sc_data:
                r = sc_data[sc]
                is_best_col = (sc == best_sc)
                p50_bg = GREEN_BG if is_best_col else bg
                _cell(ws, data_row, col,   _fmt_s(r["p50_s"]),
                      bg=p50_bg, bold=is_best_col, align="center")
                _cell(ws, data_row, col+1, _fmt_s(r["p95_s"]),
                      bg=bg, align="center")
                _cell(ws, data_row, col+2, _fmt_s(r["p99_s"]),
                      bg=bg, align="center")
            else:
                for offset in range(3):
                    _cell(ws, data_row, col+offset, "N/D", bg=LIGHT_GRAY, align="center")
            col += 3

        _cell(ws, data_row, col, best_sc or "—",
              bg=GREEN_BG if best_sc else bg, bold=True, align="center")

        data_row += 1

    data_row += 1
    sep2 = ws.cell(row=data_row, column=1,
                   value="── Variación vs PG-idx (p50) — negativo = más rápido ──────────────────")
    sep2.font = Font(bold=True, color=DARK_BLUE, size=10)
    sep2.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.merge_cells(f"A{data_row}:{get_column_letter(ncols)}{data_row}")
    data_row += 1

    if "PG-idx" in scenarios:
        _hdr(ws, data_row, 1, "Query", bg=DARK_BLUE)
        _hdr(ws, data_row, 2, "Descripción", bg=DARK_BLUE)
        col = 3
        for sc in [s for s in scenarios if s != "PG-idx"]:
            bg_h = sc_colors.get(sc, MID_BLUE)
            _hdr(ws, data_row, col, f"{sc} vs PG-idx", bg=bg_h)
            col += 1
        data_row += 1

        for qid in queries_sorted:
            sc_data = pivot[qid]
            base = sc_data.get("PG-idx", {}).get("p50_s")
            bg = LIGHT_GRAY if data_row % 2 == 0 else WHITE
            _cell(ws, data_row, 1, qid, bg=bg, bold=True, align="left")
            _cell(ws, data_row, 2, QUERY_DESCS.get(qid, ""), bg=bg, align="left")
            col = 3
            for sc in [s for s in scenarios if s != "PG-idx"]:
                if base and sc in sc_data:
                    pct = (sc_data[sc]["p50_s"] - base) / base * 100
                    txt = f"{pct:+.0f}%"
                    cell_bg = GREEN_BG if pct < -15 else (RED_BG if pct > 15 else bg)
                    _cell(ws, data_row, col, txt, bg=cell_bg, bold=abs(pct) > 15,
                          align="center")
                else:
                    _cell(ws, data_row, col, "N/D", bg=LIGHT_GRAY, align="center")
                col += 1
            data_row += 1

    _autofit(ws)
    ws.column_dimensions["B"].width = 42
    ws.row_dimensions[2].height = 36


def sheet_summary(wb: Workbook):
    ws = wb.create_sheet("Resumen ejecutivo")
    ws.sheet_view.showGridLines = False

    title = ws.cell(row=1, column=1, value="Resumen: PostgreSQL vs Apache Iceberg — Dataset DJBR (1.58M declaraciones)")
    title.font = Font(bold=True, size=13, color=DARK_BLUE)
    ws.merge_cells("A1:F1")
    title.alignment = Alignment(horizontal="center")

    ws.cell(row=3, column=1, value="CARGA (E1)").font = Font(bold=True, size=11, color=DARK_BLUE)
    ws.merge_cells("A3:F3")

    hdr_load = ["Motor", "Workers", "Throughput", "Tiempo total", "Speedup", "Filas cargadas"]
    for c, h in enumerate(hdr_load, 1):
        _hdr(ws, 4, c, h, bg=MID_BLUE)
    load_summary = [
        ("PostgreSQL", "1 (serial)",  "778 rows/s",  "33.9 min",  "baseline",  "1,584,534"),
        ("PostgreSQL", "8 (paralelo)", "2,434 rows/s", "10.9 min", "3.1x",     "1,584,534"),
        ("Iceberg",    "1 (serial)",  "917 rows/s",  "28.7 min",  "baseline",  "1,584,561"),
        ("Iceberg",    "8 (paralelo)", "3,593 rows/s", "7.4 min",  "3.9x",     "1,584,561"),
    ]
    for i, row_data in enumerate(load_summary, 5):
        bg = GREEN_BG if "8" in row_data[1] else (LIGHT_BLUE if i % 2 == 0 else WHITE)
        for c, val in enumerate(row_data, 1):
            _cell(ws, i, c, val, bg=bg, align="center",
                  bold=("3.9x" in val or "3.1x" in val))

    row_q = 11
    ws.cell(row=row_q, column=1, value="QUERIES (E3 — DJBR)").font = Font(bold=True, size=11, color=DARK_BLUE)
    ws.merge_cells(f"A{row_q}:F{row_q}")

    hdr_q = ["Caso de uso", "PG-idx p50", "PG-noidx p50", "Iceberg p50", "Ganador", "Ventaja"]
    for c, h in enumerate(hdr_q, 1):
        _hdr(ws, row_q+1, c, h, bg=MID_BLUE)

    query_summary = [
        ("Lookup por ID (PK)",               "<1ms",   "<1ms",  "549ms",  "PG-idx",  ">>1000x"),
        ("Lookup por cargo (40K filas)",      "282ms",  "602ms", "811ms",  "PG-idx",  "2.1x vs noidx"),
        ("Rango de fecha, LIMIT 500",         "2ms",    "532ms", "632ms",  "PG-idx",  "266x vs noidx"),
        ("SUM/AVG por año (GROUP BY)",        "2.06s",  "2.12s", "886ms",  "Iceberg", "2.3x"),
        ("Top 20 cargos por activo",          "1.63s",  "1.57s", "1.42s",  "Iceberg", "1.1x"),
        ("COUNT rango numérico",              "510ms",  "584ms", "731ms",  "PG-idx",  "1.4x"),
        ("COUNT(*) total",                    "470ms",  "501ms", "145ms",  "Iceberg", "3.2x"),
        ("PK lookup + JSONB extract",         "<1ms",   "<1ms",  "580ms",  "PG-idx",  ">>1000x"),
        ("GROUP BY campo JSON (full scan)",   "20.84s", "19.53s","31.16s", "PG-noidx","1.1x"),
        ("JSONB containment (@> / LIKE)",     "7.54s",  "7.70s", "20.77s", "PG-idx",  "2.8x"),
    ]
    for i, row_data in enumerate(query_summary, row_q+2):
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        winner = row_data[4]
        _cell(ws, i, 1, row_data[0], bg=bg, align="left")
        for c in range(1, 4):
            val = row_data[c]
            is_best = (c == 1 and winner == "PG-idx") or \
                      (c == 2 and winner == "PG-noidx") or \
                      (c == 3 and winner == "Iceberg")
            cell_bg = GREEN_BG if is_best else bg
            _cell(ws, i, c+1, val, bg=cell_bg, bold=is_best, align="center")
        win_bg = GREEN_BG if winner == "Iceberg" else (LIGHT_BLUE if winner == "PG-idx" else bg)
        _cell(ws, i, 5, winner, bg=win_bg, bold=True, align="center")
        _cell(ws, i, 6, row_data[5], bg=bg, align="center")

    _autofit(ws)
    ws.column_dimensions["A"].width = 40
    ws.row_dimensions[1].height = 28


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--results-dir", default=None)
    parser.add_argument("--output",      default=None)
    args = parser.parse_args()

    script_dir = Path(__file__).resolve().parent
    results_dir = Path(args.results_dir) if args.results_dir else script_dir / "results"
    out_path = Path(args.output) if args.output else \
        script_dir.parents[1] / "docs" / "experimentos" / "benchmark_resultados.xlsx"

    wb = Workbook()
    wb.remove(wb.active)

    sheet_summary(wb)

    sheet_carga(wb)

    e3_rows = merge_e3(results_dir)
    e3_scenarios = [s for s in SCENARIOS_ORDER if any(r["scenario"] == s for r in e3_rows)]
    sheet_pivot(wb, "E3 DJBR — Comparación", e3_rows, e3_scenarios)

    sheet_raw(wb, "E3 DJBR — Raw", e3_rows)

    e2_path = results_dir / "e2_2026-03-25.csv"
    if e2_path.exists():
        e2_rows = load_csv(e2_path)
        e2_scenarios = [s for s in SCENARIOS_ORDER if any(r["scenario"] == s for r in e2_rows)]
        sheet_pivot(wb, "E2 Multi-fuente — Comparación", e2_rows, e2_scenarios)
        sheet_raw(wb, "E2 Multi-fuente — Raw", e2_rows)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    print(f"Excel generado: {out_path}")


if __name__ == "__main__":
    main()
